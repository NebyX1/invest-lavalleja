import json
from pathlib import Path
import pytest
from pydantic import ValidationError
from gianna.corpus import Corpus,Source,Record,source_refs,export_jsonl,load_corpus,extract_docx
from gianna.errors import DomainError
ROOT=Path(__file__).resolve().parents[1]


def guide():
    path=ROOT/'knowledge/invest_lavalleja_2026.gianna.json'
    if not path.exists():pytest.skip('El paquete documental privado no está versionado en este checkout.')
    return Corpus.model_validate_json(path.read_text(encoding='utf-8'))


def test_guide_inventory_and_coverage():
    c=guide();report=c.conversion_report
    assert len(c.records)==271 and len(c.sources)==50
    assert sum(r.audience=='public' for r in c.records)==196
    assert sum(r.audience=='internal' for r in c.records)==75
    assert report['unaccounted_block_ids']==[]
    assert report['blocks_total']==769 and report['tables']==36
    assert all(r.cutoff_date=='2026-09-18' for r in c.records)
    assert all(s.url.startswith('https://') or s.url.startswith('http://') for s in c.sources)


def test_opportunities_intact():
    c=guide();byid={r.id:r for r in c.records}
    assert all(f'O{i:02}' in byid for i in range(1,17))
    for i in range(1,17):
        r=byid[f'O{i:02}'];assert r.evidence_type=='hypothesis'
        assert 'Condiciones.' in r.text and 'No avanzar cuando' in r.text
    assert 'padrón' in byid['O05'].text and 'hipótesis' not in byid['O05'].title.lower()


def test_financial_example_not_observed_data():
    records=[r for r in guide().records if r.topic=='EJEMPLO']
    assert records and all(r.evidence_type=='didactic_example' for r in records)
    assert all(any('hipotético' in c or 'didácticos' in c for c in r.conditions) for r in records)


def test_source_ranges():
    assert source_refs('Texto [S29–S31] y [S03, S07, S11].')==['S03','S07','S11','S29','S30','S31']


def test_jsonl_roundtrip(tmp_path,corpus):
    p=tmp_path/'test.jsonl';p.write_text(export_jsonl(corpus),encoding='utf8')
    result=load_corpus(p);assert result.model_dump()==corpus.model_dump()


def test_missing_source_rejected(corpus):
    data=corpus.model_dump();data['records'][0]['source_refs']=['missing']
    with pytest.raises(ValidationError):Corpus.model_validate(data)


def test_javascript_link_rejected():
    with pytest.raises(ValidationError):Source(id='evil',title='Evil',url='javascript:alert(1)')
    with pytest.raises(ValidationError):Source(id='evil',title='Evil',url='https://password:secret@example.test')


def test_generic_markdown_defaults_internal(tmp_path):
    p=tmp_path/'simple.md';p.write_text('# Título\nContenido de prueba.\n## Otro\nOtro contenido.',encoding='utf8')
    c=load_corpus(p);assert c.records and all(r.audience=='internal' for r in c.records)


def test_docx_linebreaks_preserved(tmp_path):
    from docx import Document
    d=Document();p=d.add_paragraph();p.add_run('Primera');p.add_run().add_break();p.add_run('Segunda')
    path=tmp_path/'test.docx';d.save(path)
    assert extract_docx(path)[0]['text']=='Primera\nSegunda'


def test_duplicate_record_rejected(corpus):
    data=corpus.model_dump();data['records'].append(data['records'][0])
    with pytest.raises(ValidationError):Corpus.model_validate(data)
